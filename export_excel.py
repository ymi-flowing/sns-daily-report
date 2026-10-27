"""
SNS Daily Report - Excel export.

Given the same data dicts that generate_html() consumes, write an .xlsx that
mirrors the dashboard grid: region-grouped property columns + Portfolio Total,
with Unit Status / Traffic / Expiring Leases / Renovation / Financial sections
stacked vertically on a single sheet.
"""
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROPERTY_GROUPS = [
    ("Tampa", ["49thStr"]),
    ("Orlando", ["TV", "MOC", "MUSO"]),
    ("Jacksonville", ["AP", "SJ", "ROP", "ORT"]),
    ("Ft. Lauderdale", ["EAST"]),
]

ALL_PROPS = [abbr for _, abbrs in PROPERTY_GROUPS for abbr in abbrs]

TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="1F3A5F")

REGION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGION_FILL = PatternFill("solid", fgColor="4472C4")

PROP_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
PROP_FILL = PatternFill("solid", fgColor="5B7FB8")

SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="2E5C8A")

METRIC_FONT = Font(name="Calibri", size=10, bold=True)
METRIC_FILL = PatternFill("solid", fgColor="EAF0F7")

SUBHDR_FONT = Font(name="Calibri", size=10, bold=True, italic=True)
SUBHDR_FILL = PatternFill("solid", fgColor="D9E4F1")

TOTAL_FONT = Font(name="Calibri", size=10, bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="F4E9C1")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN = Side(style="thin", color="B0B7BE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _write_row(ws, row, label, values, *, metric_font=METRIC_FONT,
               metric_fill=METRIC_FILL, number_format=None,
               total_value=None):
    """Write one data row: metric label, per-property values, portfolio total."""
    ws.cell(row=row, column=1, value=label).font = metric_font
    ws.cell(row=row, column=1).fill = metric_fill
    ws.cell(row=row, column=1).alignment = LEFT
    ws.cell(row=row, column=1).border = BORDER

    numeric_total = 0.0
    numeric_seen = False
    for i, v in enumerate(values, start=2):
        cell = ws.cell(row=row, column=i, value=v if v is not None else "")
        cell.alignment = CENTER
        cell.border = BORDER
        if number_format:
            cell.number_format = number_format
        if isinstance(v, (int, float)):
            numeric_total += v
            numeric_seen = True

    total_col = 2 + len(ALL_PROPS)
    if total_value is not None:
        t = total_value
    elif numeric_seen:
        t = numeric_total
    else:
        t = ""
    cell = ws.cell(row=row, column=total_col, value=t)
    cell.alignment = CENTER
    cell.border = BORDER
    cell.font = TOTAL_FONT
    cell.fill = TOTAL_FILL
    if number_format:
        cell.number_format = number_format


def _write_section_header(ws, row, title):
    total_col = 2 + len(ALL_PROPS)
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=1).alignment = LEFT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col)
    for col in range(1, total_col + 1):
        ws.cell(row=row, column=col).border = BORDER
        if col > 1:
            ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.row_dimensions[row].height = 22


def _write_subheader(ws, row, label):
    total_col = 2 + len(ALL_PROPS)
    ws.cell(row=row, column=1, value=label).font = SUBHDR_FONT
    ws.cell(row=row, column=1).fill = SUBHDR_FILL
    ws.cell(row=row, column=1).alignment = LEFT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col)
    for col in range(1, total_col + 1):
        ws.cell(row=row, column=col).border = BORDER
        if col > 1:
            ws.cell(row=row, column=col).fill = SUBHDR_FILL


def _write_header(ws, timestamp_str):
    total_col = 2 + len(ALL_PROPS)

    ws.cell(row=1, column=1, value=f"SNS Daily Report  —  {timestamp_str}")
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.cell(row=1, column=1).fill = TITLE_FILL
    ws.cell(row=1, column=1).alignment = LEFT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_col)
    for col in range(1, total_col + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    ws.row_dimensions[1].height = 26

    ws.cell(row=2, column=1, value="Metric").font = PROP_FONT
    ws.cell(row=2, column=1).fill = PROP_FILL
    ws.cell(row=2, column=1).alignment = CENTER
    ws.cell(row=2, column=1).border = BORDER

    col = 2
    for region, abbrs in PROPERTY_GROUPS:
        start = col
        end = col + len(abbrs) - 1
        ws.cell(row=2, column=start, value=region).font = REGION_FONT
        ws.cell(row=2, column=start).fill = REGION_FILL
        ws.cell(row=2, column=start).alignment = CENTER
        if end > start:
            ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        for c in range(start, end + 1):
            ws.cell(row=2, column=c).fill = REGION_FILL
            ws.cell(row=2, column=c).border = BORDER
        col = end + 1
    ws.cell(row=2, column=col, value="Portfolio").font = REGION_FONT
    ws.cell(row=2, column=col).fill = REGION_FILL
    ws.cell(row=2, column=col).alignment = CENTER
    ws.cell(row=2, column=col).border = BORDER

    ws.cell(row=3, column=1, value="").fill = PROP_FILL
    ws.cell(row=3, column=1).border = BORDER
    for i, abbr in enumerate(ALL_PROPS, start=2):
        c = ws.cell(row=3, column=i, value=abbr)
        c.font = PROP_FONT
        c.fill = PROP_FILL
        c.alignment = CENTER
        c.border = BORDER
    tot = ws.cell(row=3, column=total_col, value="Total")
    tot.font = PROP_FONT
    tot.fill = PROP_FILL
    tot.alignment = CENTER
    tot.border = BORDER

    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20


def _write_unit_status(ws, row, properties_data):
    _write_section_header(ws, row, "Unit Status")
    row += 1

    def sb(p, key):
        return int((p.get("status_breakdown") or {}).get(key, 0) or 0)

    def per(fn):
        return [fn(properties_data.get(a, {})) for a in ALL_PROPS]

    total_units = per(lambda p: int(p.get("total_units", 0) or 0))
    occupied = per(lambda p: sb(p, "Occupied"))
    vac_ready = per(lambda p: sb(p, "Vacant Ready"))
    vac_notready = per(lambda p: sb(p, "Vacant Not Ready"))
    ntv = per(lambda p: sb(p, "NTV"))
    eviction = per(lambda p: sb(p, "Under Eviction"))
    vac_leased = per(lambda p: sb(p, "Vacant Leased"))
    ntv_leased = per(lambda p: sb(p, "NTV Leased"))
    off_market = per(lambda p: int(p.get("off_market", 0) or 0))

    def phys_occ(p):
        t = int(p.get("total_units", 0) or 0)
        if not t:
            return 0.0
        vac = sb(p, "Vacant Ready") + sb(p, "Vacant Not Ready")
        return round((t - vac) / t * 100, 2)

    def occ_rented(p):
        t = int(p.get("total_units", 0) or 0)
        if not t:
            return 0.0
        vac = sb(p, "Vacant Ready") + sb(p, "Vacant Not Ready")
        leased = sb(p, "Vacant Leased")
        return round((t - vac + leased) / t * 100, 2)

    def rentable(p):
        return (sb(p, "Vacant Ready") + sb(p, "Vacant Not Ready")
                + sb(p, "NTV") - sb(p, "Vacant Leased") - sb(p, "NTV Leased"))

    phys = per(phys_occ)
    rented = per(occ_rented)
    rentable_units = per(rentable)

    tt = sum(total_units) or 1
    portfolio_phys = round((tt - sum(vac_ready) - sum(vac_notready)) / tt * 100, 2)
    portfolio_rented = round(
        (tt - sum(vac_ready) - sum(vac_notready) + sum(vac_leased)) / tt * 100, 2
    )

    rows = [
        ("Total Units", total_units, None, None),
        ("Occupied", occupied, None, None),
        ("Vacant Ready", vac_ready, None, None),
        ("Vacant Not Ready", vac_notready, None, None),
        ("NTV", ntv, None, None),
        ("Under Eviction", eviction, None, None),
        ("Physical Occupancy %", phys, "0.00", portfolio_phys),
        ("Vacant Leased", vac_leased, None, None),
        ("NTV Leased", ntv_leased, None, None),
        ("Rentable Units", rentable_units, None, None),
        ("Off the Market", off_market, None, None),
        ("Occupancy Rented %", rented, "0.00", portfolio_rented),
    ]
    for label, values, fmt, total in rows:
        _write_row(ws, row, label, values, number_format=fmt, total_value=total)
        row += 1
    return row + 1


def _write_traffic(ws, row, traffic_data, leases_data):
    def t(a, key):
        return int((traffic_data.get(a) or {}).get(key, 0) or 0)

    def tf(a, key):
        try:
            return float((traffic_data.get(a) or {}).get(key, 0) or 0)
        except (TypeError, ValueError):
            return 0.0

    def l(a, key):
        return int((leases_data.get(a) or {}).get(key, 0) or 0)

    def conv_pct(a):
        prospects = t(a, "total")
        signed = l(a, "signed_leases_mtd")
        return round((signed / prospects) * 100, 1) if prospects > 0 else 0.0

    _write_section_header(ws, row, "Traffic (MTD)")
    row += 1
    mtd_rows = [
        ("Total Prospects (MTD)", [t(a, "total") for a in ALL_PROPS], None, None),
        ("Daily Avg Prospects", [round(tf(a, "daily_avg"), 1) for a in ALL_PROPS], "0.0", None),
        ("New Signed Leases (MTD)", [l(a, "signed_leases_mtd") for a in ALL_PROPS], None, None),
        ("Conversion Rate %", [conv_pct(a) for a in ALL_PROPS], "0.0", None),
        ("Applicants MTD", [l(a, "applicants_mtd") for a in ALL_PROPS], None, None),
        ("Approved (MTD)", [l(a, "applicants_mtd_approved") for a in ALL_PROPS], None, None),
        ("Denied (MTD)", [l(a, "applicants_mtd_denied") for a in ALL_PROPS], None, None),
        ("Cancelled (MTD)", [l(a, "applicants_mtd_cancelled") for a in ALL_PROPS], None, None),
        ("Still Active (MTD)", [l(a, "applicants_mtd_active") for a in ALL_PROPS], None, None),
    ]
    for label, values, fmt, total in mtd_rows:
        _write_row(ws, row, label, values, number_format=fmt, total_value=total)
        row += 1
    row += 1

    _write_section_header(ws, row, "Traffic (Snapshot)")
    row += 1
    snapshot_rows = [
        ("Total Applicants", [l(a, "total_applicants") for a in ALL_PROPS], None, None),
        ("Pending Approval", [l(a, "pending_approval") for a in ALL_PROPS], None, None),
        ("Approved", [l(a, "approved") for a in ALL_PROPS], None, None),
        ("Signed Lease", [l(a, "signed_lease") for a in ALL_PROPS], None, None),
    ]
    for label, values, fmt, total in snapshot_rows:
        _write_row(ws, row, label, values, number_format=fmt, total_value=total)
        row += 1
    return row + 1


def _write_expiring(ws, row, leases_data):
    _write_section_header(ws, row, "Expiring Leases (Current + Next 3 Months)")
    row += 1

    def exp(a):
        return ((leases_data.get(a) or {}).get("expiring") or {})

    labels = {}
    for a in ALL_PROPS:
        e = exp(a)
        if e.get("month1_label"):
            labels = {
                "month1": e.get("month1_label"),
                "month2": e.get("month2_label"),
                "month3": e.get("month3_label"),
                "month4": e.get("month4_label"),
            }
            break
    if not labels:
        labels = {"month1": "M1", "month2": "M2", "month3": "M3", "month4": "M4"}

    statuses = ["Current", "Pending Renewal", "Renewed", "NTV/Formers", "Under Eviction"]

    def sc(a, month_key, status):
        return int((exp(a).get(month_key) or {}).get(status, 0) or 0)

    def mt(a, month_key):
        m = exp(a).get(month_key) or {}
        return sum(int(v or 0) for v in m.values())

    def renewal_pct(a, month_key):
        m = exp(a).get(month_key) or {}
        renewed = int(m.get("Renewed", 0) or 0)
        denom = (renewed
                 + int(m.get("Current", 0) or 0)
                 + int(m.get("NTV/Formers", 0) or 0)
                 + int(m.get("Under Eviction", 0) or 0))
        if denom == 0:
            return 0.0
        return round(renewed / denom * 100, 1)

    def p_status(month_key, status):
        return sum(sc(a, month_key, status) for a in ALL_PROPS)

    def p_total(month_key):
        return sum(mt(a, month_key) for a in ALL_PROPS)

    def p_renewal(month_key):
        renewed = p_status(month_key, "Renewed")
        denom = (renewed
                 + p_status(month_key, "Current")
                 + p_status(month_key, "NTV/Formers")
                 + p_status(month_key, "Under Eviction"))
        if denom == 0:
            return 0.0
        return round(renewed / denom * 100, 1)

    for month_key in ("month1", "month2", "month3", "month4"):
        _write_subheader(ws, row, labels[month_key])
        row += 1
        for status in statuses:
            values = [sc(a, month_key, status) for a in ALL_PROPS]
            _write_row(ws, row, status, values,
                       total_value=p_status(month_key, status))
            row += 1
        totals = [mt(a, month_key) for a in ALL_PROPS]
        _write_row(ws, row, "Total Expiring", totals,
                   total_value=p_total(month_key))
        row += 1
        pcts = [renewal_pct(a, month_key) for a in ALL_PROPS]
        _write_row(ws, row, "Renewal %", pcts, number_format="0.0",
                   total_value=p_renewal(month_key))
        row += 1

    _write_subheader(ws, row, "4-Month Total")
    row += 1
    for status in statuses:
        values = [int((exp(a).get("total") or {}).get(status, 0) or 0) for a in ALL_PROPS]
        _write_row(ws, row, status, values)
        row += 1

    return row + 1


def _write_renovation(ws, row, leases_data):
    _write_section_header(ws, row, "Renovation Status")
    row += 1

    def rn(a, key):
        return int(((leases_data.get(a) or {}).get("renovation") or {}).get(key, 0) or 0)

    rows = [
        ("Renovated", [rn(a, "renovated") for a in ALL_PROPS]),
        ("Partial Renovated", [rn(a, "partial") for a in ALL_PROPS]),
        ("Classic", [rn(a, "classic") for a in ALL_PROPS]),
        ("Total Units", [rn(a, "total") for a in ALL_PROPS]),
    ]
    for label, values in rows:
        _write_row(ws, row, label, values)
        row += 1
    return row + 1


def _write_work_orders(ws, row, leases_data):
    _write_section_header(ws, row, "Work Orders")
    row += 1

    def wo(a, key):
        return int(((leases_data.get(a) or {}).get("work_orders") or {}).get(key, 0) or 0)

    _write_row(ws, row, "Open (non-make-ready)", [wo(a, "open_total") for a in ALL_PROPS]); row += 1
    _write_row(ws, row, "New in Last 24h", [wo(a, "new_24h") for a in ALL_PROPS]); row += 1
    return row + 1


def _write_future_moves(ws, row, leases_data):
    _write_section_header(ws, row, "Future Moves (Current Month)")
    row += 1

    def fm(a, key):
        return int(((leases_data.get(a) or {}).get("future_moves") or {}).get(key, 0) or 0)

    _write_row(ws, row, "Future Move-Outs", [fm(a, "future_mo") for a in ALL_PROPS]); row += 1
    return row + 1


def _write_financial(ws, row, financial_data):
    _write_section_header(ws, row, "Financial Summary")
    row += 1

    income = financial_data.get("income") or {}
    delinq = financial_data.get("delinquency") or {}

    def inc(a, key):
        return _num((income.get(a) or {}).get(key, 0))

    def dg(a, cat, key):
        return _num(((delinq.get(a) or {}).get(cat) or {}).get(key, 0))

    charges = [inc(a, "charge") for a in ALL_PROPS]
    credits = [inc(a, "credit") for a in ALL_PROPS]
    net = [c - k for c, k in zip(charges, credits)]
    rec_mtd = [inc(a, "receivable_mtd") for a in ALL_PROPS]

    def del_pct(net_v, rec_v):
        return round(rec_v / net_v * 100, 2) if net_v else 0.0

    del_pct_list = [del_pct(n, r) for n, r in zip(net, rec_mtd)]

    _write_subheader(ws, row, "Income"); row += 1
    _write_row(ws, row, "Total Income", charges, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Total Credit", credits, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Net", net, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Receivable MTD", rec_mtd, number_format="$#,##0.00"); row += 1

    portfolio_net = sum(net)
    portfolio_rec = sum(rec_mtd)
    portfolio_del_pct = round(portfolio_rec / portfolio_net * 100, 2) if portfolio_net else 0.0
    _write_row(ws, row, "Del %", del_pct_list, number_format="0.00",
               total_value=portfolio_del_pct); row += 1

    _write_subheader(ws, row, "Delinquency (amount / resident count)"); row += 1

    current_amt = [dg(a, "current", "amount") for a in ALL_PROPS]
    current_cnt = [int(dg(a, "current", "count")) for a in ALL_PROPS]
    evic_amt = [dg(a, "eviction", "amount") for a in ALL_PROPS]
    evic_cnt = [int(dg(a, "eviction", "count")) for a in ALL_PROPS]
    former_amt = [dg(a, "formers", "amount") for a in ALL_PROPS]
    former_cnt = [int(dg(a, "formers", "count")) for a in ALL_PROPS]
    total_amt = [a + b + c for a, b, c in zip(current_amt, evic_amt, former_amt)]
    total_cnt = [a + b + c for a, b, c in zip(current_cnt, evic_cnt, former_cnt)]

    _write_row(ws, row, "Current  ($)", current_amt, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Current  (#)", current_cnt); row += 1
    _write_row(ws, row, "Under Eviction  ($)", evic_amt, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Under Eviction  (#)", evic_cnt); row += 1
    _write_row(ws, row, "Formers  ($)", former_amt, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Formers  (#)", former_cnt); row += 1
    _write_row(ws, row, "Total Delinquency  ($)", total_amt, number_format="$#,##0.00"); row += 1
    _write_row(ws, row, "Total Delinquency  (#)", total_cnt); row += 1

    return row + 1


def export_dashboard_xlsx(properties_data, traffic_data, leases_data,
                          financial_data, output_path):
    """Write the dashboard workbook to output_path (overwrites)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    florida_tz = ZoneInfo("America/New_York")
    now = datetime.now(florida_tz)
    timestamp_str = now.strftime("%B %d, %Y at %I:%M %p ET")

    _write_header(ws, timestamp_str)

    row = 5
    row = _write_unit_status(ws, row, properties_data)
    row = _write_traffic(ws, row, traffic_data, leases_data)
    row = _write_expiring(ws, row, leases_data)
    row = _write_renovation(ws, row, leases_data)
    row = _write_work_orders(ws, row, leases_data)
    row = _write_future_moves(ws, row, leases_data)
    row = _write_financial(ws, row, financial_data)

    ws.column_dimensions["A"].width = 32
    total_col = 2 + len(ALL_PROPS)
    for i in range(2, total_col):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.column_dimensions[get_column_letter(total_col)].width = 13

    ws.freeze_panes = "B4"

    wb.save(output_path)
    return output_path
